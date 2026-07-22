from __future__ import annotations

import os
import shutil
import sqlite3
import tempfile
from contextlib import closing
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parent
DEFAULT_DB_PATH = Path(os.getenv("AL_SADI_DB_PATH", ROOT / "factory.db"))
WORKBOOK_PATH = Path(os.getenv("AL_SADI_WORKBOOK_PATH", ROOT / "factory_records.xlsx"))
EXPORTS_PATH = Path(os.getenv("AL_SADI_EXPORTS_PATH", ROOT / "exports"))
BACKUPS_PATH = Path(os.getenv("AL_SADI_BACKUPS_PATH", ROOT / "backups"))
SHEET_PASSWORD = os.getenv("AL_SADI_SHEET_PASSWORD", "alsadi-local-protection")
REQUIRED_SHEETS = ["Production", "Expenses", "Employees", "Attendance", "Machines", "Summary"]


def _read_sql(query: str, conn: sqlite3.Connection) -> pd.DataFrame:
    return pd.read_sql_query(query, conn)


def _currency(value: float) -> float:
    return round(float(value or 0), 2)


def _load_frames(db_path: Path) -> dict[str, pd.DataFrame]:
    with closing(sqlite3.connect(db_path)) as conn:
        return {
            "production": _read_sql("SELECT * FROM production_entries ORDER BY production_date, id", conn),
            "expenses": _read_sql("SELECT * FROM expenses ORDER BY expense_date, id", conn),
            "employees": _read_sql(
                """
                SELECT id, name, role, phone, salary, advance, performance_score, status,
                       joined_on, archived_at, updated_at
                FROM employees ORDER BY name, id
                """,
                conn,
            ),
            "attendance": _read_sql(
                """
                SELECT a.id, a.employee_id, e.name AS employee_name, a.attendance_date,
                       a.status, a.notes, a.created_at, a.updated_at
                FROM attendance a JOIN employees e ON e.id = a.employee_id
                ORDER BY a.attendance_date, a.id
                """,
                conn,
            ),
            "machines": _read_sql("SELECT * FROM machines ORDER BY machine_number, id", conn),
        }


def build_summary(
    production: pd.DataFrame,
    expenses: pd.DataFrame,
    machines: pd.DataFrame,
) -> dict[str, pd.DataFrame]:
    total_earnings = _currency(production["total_amount"].sum()) if not production.empty else 0
    total_expenses = _currency(expenses["amount"].sum()) if not expenses.empty else 0
    kpis = pd.DataFrame(
        [
            {"Metric": "Total Earnings", "Value": total_earnings},
            {"Metric": "Total Expenses", "Value": total_expenses},
            {"Metric": "Profit", "Value": _currency(total_earnings - total_expenses)},
            {"Metric": "Machines", "Value": len(machines)},
        ]
    )

    if production.empty:
        machine_wise = pd.DataFrame(columns=["Machine Number", "Total Quantity", "Total Earnings"])
        monthly_income = pd.DataFrame(columns=["Month", "Total Quantity", "Earnings"])
    else:
        machine_wise = (
            production.groupby("machine_number", as_index=False)
            .agg(total_quantity=("quantity", "sum"), total_earnings=("total_amount", "sum"))
            .rename(
                columns={
                    "machine_number": "Machine Number",
                    "total_quantity": "Total Quantity",
                    "total_earnings": "Total Earnings",
                }
            )
        )
        production_dates = pd.to_datetime(production["production_date"])
        monthly_income = (
            production.assign(month=production_dates.dt.to_period("M").astype(str))
            .groupby("month", as_index=False)
            .agg(total_quantity=("quantity", "sum"), earnings=("total_amount", "sum"))
            .rename(columns={"month": "Month", "total_quantity": "Total Quantity", "earnings": "Earnings"})
        )

    if expenses.empty:
        monthly_expenses = pd.DataFrame(columns=["Month", "Expenses"])
    else:
        expense_dates = pd.to_datetime(expenses["expense_date"])
        monthly_expenses = (
            expenses.assign(month=expense_dates.dt.to_period("M").astype(str))
            .groupby("month", as_index=False)
            .agg(expenses=("amount", "sum"))
            .rename(columns={"month": "Month", "expenses": "Expenses"})
        )

    monthly_totals = monthly_income.merge(monthly_expenses, on="Month", how="outer").fillna(0)
    if monthly_totals.empty:
        monthly_totals = pd.DataFrame(
            columns=["Month", "Total Quantity", "Earnings", "Expenses", "Profit"]
        )
    else:
        monthly_totals["Profit"] = monthly_totals["Earnings"] - monthly_totals["Expenses"]
        monthly_totals = monthly_totals.sort_values("Month")
    return {"kpis": kpis, "machine_wise": machine_wise, "monthly_totals": monthly_totals}


def _format_worksheet(writer: pd.ExcelWriter, sheet_name: str) -> None:
    ws = writer.book[sheet_name]
    header_fill = PatternFill("solid", fgColor="0B5D7A")
    header_font = Font(color="FFFFFF", bold=True)
    title_font = Font(color="DDF8FF", bold=True, size=12)

    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="center", wrap_text=False)

    header_rows = [1]
    if sheet_name == "Summary":
        header_rows = [1, 7, 14]
    for row_number in header_rows:
        for cell in ws[row_number]:
            if cell.value is not None:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(row_number, 1).font = title_font

    date_headers = {"production_date", "expense_date", "joined_on", "attendance_date", "installed_on"}
    money_headers = {"rate_per_unit", "total_amount", "amount", "salary", "advance", "Value", "Earnings", "Expenses", "Profit", "Total Earnings"}
    if sheet_name != "Summary":
        header_map = {cell.value: cell.column for cell in ws[1]}
        for header in date_headers & set(header_map):
            for cell in ws.iter_cols(min_col=header_map[header], max_col=header_map[header], min_row=2):
                for item in cell:
                    item.number_format = "yyyy-mm-dd"
        for header in money_headers & set(header_map):
            for cell in ws.iter_cols(min_col=header_map[header], max_col=header_map[header], min_row=2):
                for item in cell:
                    item.number_format = '#,##0.00'

    ws.freeze_panes = "A2"
    ws.sheet_view.showGridLines = False
    if ws.max_row and ws.max_column and sheet_name != "Summary":
        ws.auto_filter.ref = ws.dimensions
    for column_cells in ws.columns:
        letter = get_column_letter(column_cells[0].column)
        max_length = max(len(str(cell.value or "")) for cell in column_cells)
        ws.column_dimensions[letter].width = min(max(max_length + 3, 13), 36)


def _protect_worksheet(writer: pd.ExcelWriter, sheet_name: str) -> None:
    ws = writer.book[sheet_name]
    for row in ws.iter_rows():
        for cell in row:
            cell.protection = Protection(locked=True)
    if sheet_name in {"Production", "Expenses"}:
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.protection = Protection(locked=False)

    ws.protection.sheet = True
    ws.protection.password = SHEET_PASSWORD
    ws.protection.selectLockedCells = False
    ws.protection.selectUnlockedCells = True
    ws.protection.formatCells = False
    ws.protection.formatColumns = False
    ws.protection.formatRows = False
    ws.protection.insertRows = False
    ws.protection.deleteRows = False
    ws.protection.sort = False
    ws.protection.autoFilter = False


def _write_workbook_file(path: Path, frames: dict[str, pd.DataFrame]) -> None:
    summary = build_summary(frames["production"], frames["expenses"], frames["machines"])
    with pd.ExcelWriter(path, engine="openpyxl", mode="w") as writer:
        frames["production"].to_excel(writer, sheet_name="Production", index=False)
        frames["expenses"].to_excel(writer, sheet_name="Expenses", index=False)
        frames["employees"].to_excel(writer, sheet_name="Employees", index=False)
        frames["attendance"].to_excel(writer, sheet_name="Attendance", index=False)
        frames["machines"].to_excel(writer, sheet_name="Machines", index=False)
        summary["kpis"].to_excel(writer, sheet_name="Summary", index=False, startrow=0)
        summary["machine_wise"].to_excel(writer, sheet_name="Summary", index=False, startrow=6)
        summary["monthly_totals"].to_excel(writer, sheet_name="Summary", index=False, startrow=13)
        writer.book.calculation.fullCalcOnLoad = True
        writer.book.calculation.forceFullCalc = True
        for sheet_name in REQUIRED_SHEETS:
            _format_worksheet(writer, sheet_name)
            _protect_worksheet(writer, sheet_name)


def _atomic_write(path: Path, frames: dict[str, pd.DataFrame]) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    temp_path: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(
            dir=path.parent, prefix=f".{path.stem}_", suffix=".tmp.xlsx", delete=False
        ) as handle:
            temp_path = Path(handle.name)
        _write_workbook_file(temp_path, frames)
        os.replace(temp_path, path)
        return path
    finally:
        if temp_path and temp_path.exists():
            temp_path.unlink(missing_ok=True)


def sync_factory_workbook(
    path: str | Path | None = None,
    db_path: str | Path | None = None,
) -> Path:
    destination = Path(path) if path is not None else WORKBOOK_PATH
    frames = _load_frames(Path(db_path) if db_path is not None else DEFAULT_DB_PATH)
    return _atomic_write(destination, frames)


def export_month_workbook(
    month: str,
    path: str | Path | None = None,
    db_path: str | Path | None = None,
) -> Path:
    try:
        pd.Period(month, freq="M")
    except ValueError as exc:
        raise ValueError("Month must use YYYY-MM format.") from exc

    frames = _load_frames(Path(db_path) if db_path is not None else DEFAULT_DB_PATH)
    for key, date_column in {
        "production": "production_date",
        "expenses": "expense_date",
        "attendance": "attendance_date",
    }.items():
        frame = frames[key]
        if not frame.empty:
            mask = pd.to_datetime(frame[date_column]).dt.to_period("M").astype(str) == month
            frames[key] = frame.loc[mask].copy()

    destination = Path(path) if path is not None else EXPORTS_PATH / f"factory_records_{month}.xlsx"
    return _atomic_write(destination, frames)


def _unique_backup_path(directory: Path, stem: str, suffix: str) -> Path:
    directory.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
    candidate = directory / f"{stem}_{timestamp}{suffix}"
    index = 1
    while candidate.exists():
        candidate = directory / f"{stem}_{timestamp}_{index}{suffix}"
        index += 1
    return candidate


def backup_excel_file(
    source_file: str | Path | None = None,
    backup_dir: str | Path | None = None,
) -> tuple[str | None, str]:
    source_path = Path(source_file) if source_file is not None else WORKBOOK_PATH
    if not source_path.exists():
        return None, "Excel file not found."
    destination = _unique_backup_path(
        Path(backup_dir) if backup_dir is not None else BACKUPS_PATH,
        "factory_records_backup",
        ".xlsx",
    )
    shutil.copy2(source_path, destination)
    return str(destination), "Backup created successfully."
