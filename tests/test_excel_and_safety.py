from __future__ import annotations

import sqlite3
from datetime import date
from pathlib import Path

import openpyxl
import pytest

import database
import spreadsheet_sync
from data_safety import backup_database, restore_database, validate_sqlite_database
from database import ValidationError


def _seed_two_months(db_path: Path, actor) -> None:
    employee = database.create_employee(
        "Excel Worker", "Operator", "", 1000, 0, 80, "Active", date(2026, 1, 1), actor, db_path
    )
    database.save_machine("M-01", "Running", "Excel Worker", "", "2026-01-01", actor, db_path)
    database.create_production("2026-06-30", "M-01", "Excel Worker", "Cuff", 10, 2, actor, db_path)
    database.create_production("2026-07-01", "M-01", "Excel Worker", "Collar", 20, 3, actor, db_path)
    database.create_expense("Power", 5, "June", "2026-06-30", actor, db_path)
    database.create_expense("Yarn", 7, "July", "2026-07-02", actor, db_path)
    database.create_attendance(employee.entity_id, "2026-07-01", "Present", "", actor, db_path)


def test_excel_has_required_protected_sheets(isolated_factory, actor, tmp_path):
    db_path, _ = isolated_factory
    _seed_two_months(db_path, actor)
    output = spreadsheet_sync.sync_factory_workbook(tmp_path / "checked.xlsx", db_path)
    workbook = openpyxl.load_workbook(output)
    assert workbook.sheetnames == spreadsheet_sync.REQUIRED_SHEETS
    for sheet in workbook.worksheets:
        assert sheet.protection.sheet is True
    assert workbook["Attendance"]["A2"].protection.locked is True
    assert workbook["Summary"]["B2"].protection.locked is True
    assert workbook["Production"]["A2"].protection.locked is False
    assert workbook["Expenses"]["A2"].protection.locked is False


def test_month_export_filters_raw_month_data(isolated_factory, actor, tmp_path):
    db_path, _ = isolated_factory
    _seed_two_months(db_path, actor)
    output = spreadsheet_sync.export_month_workbook("2026-07", tmp_path / "july.xlsx", db_path)
    workbook = openpyxl.load_workbook(output, data_only=False)
    production_dates = [cell.value for cell in workbook["Production"]["B"][1:] if cell.value]
    expense_dates = [cell.value for cell in workbook["Expenses"]["E"][1:] if cell.value]
    assert production_dates == ["2026-07-01"]
    assert expense_dates == ["2026-07-02"]
    assert workbook["Attendance"].max_row == 2
    assert all(sheet.protection.sheet for sheet in workbook.worksheets)


def test_excel_backup_names_are_unique(isolated_factory, actor, tmp_path):
    db_path, _ = isolated_factory
    source = spreadsheet_sync.sync_factory_workbook(tmp_path / "source.xlsx", db_path)
    backup_dir = tmp_path / "excel_backups"
    first, first_message = spreadsheet_sync.backup_excel_file(source, backup_dir)
    second, second_message = spreadsheet_sync.backup_excel_file(source, backup_dir)
    assert first_message == second_message == "Backup created successfully."
    assert first != second
    assert Path(first).exists() and Path(second).exists()


def test_backup_names_are_unique_and_integrity_passes(isolated_factory, actor, tmp_path):
    db_path, _ = isolated_factory
    backup_dir = tmp_path / "backups"
    first = backup_database(actor, db_path, backup_dir)
    second = backup_database(actor, db_path, backup_dir)
    assert first != second
    assert first.exists() and second.exists()
    assert validate_sqlite_database(first) == (True, "SQLite integrity check passed.")


def test_restore_rejects_invalid_database(isolated_factory, actor, tmp_path):
    db_path, _ = isolated_factory
    backup_dir = tmp_path / "backups"
    backup_dir.mkdir()
    invalid = backup_dir / "invalid.db"
    invalid.write_text("not a sqlite database", encoding="utf-8")
    with pytest.raises(ValidationError):
        restore_database(invalid, "RESTORE", actor, db_path, backup_dir)


def test_database_backup_can_restore_disposable_copy(isolated_factory, actor, tmp_path):
    db_path, _ = isolated_factory
    backup_dir = tmp_path / "backups"
    database.create_expense("Original", 25, "", "2026-07-01", actor, db_path)
    backup = backup_database(actor, db_path, backup_dir)
    database.create_expense("Later", 30, "", "2026-07-02", actor, db_path)
    restore_database(backup, "RESTORE", actor, db_path, backup_dir)
    with sqlite3.connect(db_path) as conn:
        assert conn.execute("SELECT COUNT(*) FROM expenses").fetchone()[0] == 1
